"""Policy-bound file and document operations."""

import io
import os
import tempfile
import zipfile
from pathlib import Path
from typing import Any


class FileTools:
    def __init__(
        self, roots: list[Path], artifacts: Path, max_bytes: int, max_files: int
    ):
        self.roots = roots
        self.artifacts = artifacts
        self.max_bytes = max_bytes
        self.max_files = max_files

    def _read_path(self, path: str) -> Path:
        if not self.roots:
            raise PermissionError("file reads are disabled: configure read_roots")
        candidate = Path(path)
        if not candidate.is_absolute():
            if len(self.roots) != 1:
                raise ValueError(
                    "use an absolute path when multiple read roots are configured"
                )
            candidate = self.roots[0] / candidate
        candidate = candidate.resolve(strict=True)
        if not any(candidate.is_relative_to(root) for root in self.roots):
            raise PermissionError("path is outside configured read_roots")
        return candidate

    def _bytes(self, path: Path) -> bytes:
        if not path.is_file():
            raise ValueError("path must be a regular file")
        with path.open("rb") as stream:
            data = stream.read(self.max_bytes + 1)
        if len(data) > self.max_bytes:
            raise ValueError("file exceeds max_file_bytes")
        return data

    def read_file(self, path: str) -> dict[str, Any]:
        """Read a UTF-8 text file inside configured read roots."""
        source = self._read_path(path)
        data = self._bytes(source)
        return {"source": str(source), "text": data.decode("utf-8"), "bytes": len(data)}

    def list_files(self, path: str = ".", recursive: bool = False) -> dict[str, Any]:
        """List bounded regular files; directory symlinks are not followed."""
        source = self._read_path(path)
        if not source.is_dir():
            raise ValueError("path must be a directory")
        files: list[str] = []
        visited = 0
        pending = [source]
        while pending:
            with os.scandir(pending.pop()) as entries:
                for entry in entries:
                    visited += 1
                    if visited > self.max_files:
                        raise ValueError("directory exceeds max_list_entries")
                    if entry.is_symlink():
                        continue
                    item = Path(entry.path)
                    if entry.is_file(follow_symlinks=False):
                        files.append(str(item.relative_to(source)))
                    elif recursive and entry.is_dir(follow_symlinks=False):
                        pending.append(item)
        files.sort()
        return {"source": str(source), "files": files}

    def parse_document(self, path: str) -> dict[str, Any]:
        """Extract UTF-8 text or PDF/DOCX pages/paragraphs; scanned PDFs need OCR elsewhere."""
        source = self._read_path(path)
        data = self._bytes(source)
        suffix = source.suffix.lower()
        pages: list[dict[str, Any]]
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError(
                    "install requirements-documents.txt for PDF parsing"
                ) from exc
            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > 500:
                raise ValueError("PDF exceeds 500 pages")
            pages = [
                {"page": index + 1, "text": page.extract_text() or ""}
                for index, page in enumerate(reader.pages)
            ]
        elif suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError(
                    "install requirements-documents.txt for DOCX parsing"
                ) from exc
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > self.max_bytes * 10
                ):
                    raise ValueError("DOCX decompressed content exceeds limit")
            document = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)
            tables = [
                "\n".join("\t".join(c.text for c in row.cells) for row in table.rows)
                for table in document.tables
            ]
            pages = [{"page": None, "text": "\n".join([text, *tables])}]
        elif suffix in {".txt", ".md", ".json", ".yaml", ".yml"}:
            pages = [{"page": None, "text": data.decode("utf-8")}]
        else:
            raise ValueError("supported formats: pdf, docx, txt, md, json, yaml, yml")
        if sum(len(page["text"].encode("utf-8")) for page in pages) > self.max_bytes:
            raise ValueError("extracted text exceeds max_file_bytes")
        return {"source": str(source), "format": suffix[1:], "pages": pages}

    def write_artifact(
        self, path: str, content: str, overwrite: bool = False
    ) -> dict[str, Any]:
        """Save UTF-8 content under artifact_dir, without overwriting by default."""
        relative = Path(path)
        if relative.is_absolute() or not relative.parts or ".." in relative.parts:
            raise ValueError("artifact path must be relative without '..'")
        target = (self.artifacts / relative).resolve()
        if not target.is_relative_to(self.artifacts):
            raise PermissionError("artifact path escapes artifact_dir")
        data = content.encode("utf-8")
        if len(data) > self.max_bytes:
            raise ValueError("artifact exceeds max_file_bytes")
        target.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=target.parent, delete=False) as stream:
            temporary = Path(stream.name)
            try:
                stream.write(data)
                stream.flush()
                os.fsync(stream.fileno())
                if overwrite:
                    os.replace(temporary, target)
                else:
                    os.link(temporary, target)
            finally:
                temporary.unlink(missing_ok=True)
        return {"path": str(target), "bytes": len(data)}
